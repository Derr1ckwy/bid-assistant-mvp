import hashlib
from pathlib import Path

from docx import Document

from bid_assistant.docx_quality import build_docx_quality_report, verify_docx_output
from bid_assistant.exporter import export_docx
from bid_assistant.models import (
    ChapterDraft,
    ProjectInfo,
    RequirementItem,
    ReviewReport,
    ScoringItem,
    TenderAnalysis,
)


def _export_test_document(path: Path) -> Path:
    analysis = TenderAnalysis(
        project_info=ProjectInfo(
            project_name="智慧园区项目",
            purchaser="测试采购单位",
            budget="180 万元",
            bid_deadline="2026-09-15 09:30",
        ),
        mandatory_requirements=[
            RequirementItem(content="必须提交营业执照。", source_page=3, status="已确认")
        ],
        scoring_items=[
            ScoringItem(
                criterion="技术方案完整性",
                points="20",
                source_page=8,
                status="已确认",
            )
        ],
    )
    drafts = [
        ChapterDraft(
            chapter_id="chapter_1",
            title="技术方案",
            markdown="## 总体设计\n\n1. 建立统一平台。\n\n- 支持人工复核。",
        )
    ]
    return export_docx(path, analysis, drafts, ReviewReport())


def test_exported_word_passes_integrity_and_layout_quality_checks(tmp_path: Path) -> None:
    target = _export_test_document(tmp_path / "投标文件_V001.docx")
    checksum = hashlib.sha256(target.read_bytes()).hexdigest()

    quality = verify_docx_output(
        target,
        expected_sha256=checksum,
        expected_chapter_count=1,
    )
    report = build_docx_quality_report(
        quality,
        word_version={"version": 1, "note": "第一次内部评审"},
    )

    assert quality["valid"] is True
    assert quality["errors"] == []
    assert quality["warnings"] == []
    assert quality["chapter_count"] == 1
    assert quality["paragraph_count"] > 10
    assert quality["table_count"] >= 2
    assert all(item["status"] == "pass" for item in quality["checks"])
    assert report.startswith(b"\xef\xbb\xbf")
    report_text = report.decode("utf-8-sig")
    assert "质检结论：通过" in report_text
    assert "Word 版本：V001" in report_text
    assert "[通过] 正文样式" in report_text


def test_word_quality_detects_version_hash_mismatch(tmp_path: Path) -> None:
    target = _export_test_document(tmp_path / "投标文件_V001.docx")

    quality = verify_docx_output(target, expected_sha256="0" * 64, expected_chapter_count=1)

    assert quality["valid"] is False
    assert "Word 文件 SHA-256 与版本记录不一致。" in quality["errors"]


def test_word_quality_rejects_corrupted_docx(tmp_path: Path) -> None:
    target = tmp_path / "损坏文件.docx"
    target.write_bytes(b"not-a-docx")

    quality = verify_docx_output(target)

    assert quality["valid"] is False
    assert "文件不是有效的 DOCX 压缩包。" in quality["errors"]


def test_word_quality_reports_missing_file_without_raising(tmp_path: Path) -> None:
    quality = verify_docx_output(tmp_path / "已删除文件.docx", expected_sha256="0" * 64)

    assert quality["valid"] is False
    assert quality["size"] == 0
    assert quality["sha256"] == ""
    assert quality["errors"] == ["Word 文件不存在。"]


def test_word_quality_reports_missing_chapters_and_placeholders(tmp_path: Path) -> None:
    target = tmp_path / "人工文件.docx"
    document = Document()
    document.add_paragraph("技术参数待补充")
    document.save(target)
    checksum = hashlib.sha256(target.read_bytes()).hexdigest()

    quality = verify_docx_output(
        target,
        expected_sha256=checksum,
        expected_chapter_count=1,
    )

    assert quality["valid"] is False
    assert quality["chapter_count"] == 0
    assert quality["placeholder_count"] == 1
    assert any("少于版本记录" in item for item in quality["errors"])
    assert any("不是 A4" in item for item in quality["errors"])
    assert any("占位内容" in item for item in quality["warnings"])
